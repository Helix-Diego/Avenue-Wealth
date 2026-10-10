import csv
import io
import re
from datetime import datetime, timedelta
from html.parser import HTMLParser

# Known Accounts
ACCOUNT_SANTANDER_MAIN = "Santander (*3012)"
ACCOUNT_SANTANDER_SAVINGS = "Santander Ahorro (*1842)"
ACCOUNT_SANTANDER_CARD = "Santander Tarjeta (*0871)"
ACCOUNT_IBERIA_CARDS = "Iberia Cards (*2366)"
ACCOUNT_N26 = "N26 (*7089)"
ACCOUNT_AMEX = "American Express Platinum"
ACCOUNT_MYINVESTOR = "MyInvestor"

SPANISH_MONTH_MAP = {
    'ene': '01', 'feb': '02', 'mar': '03', 'abr': '04',
    'may': '05', 'jun': '06', 'jul': '07', 'ago': '08',
    'sep': '09', 'oct': '10', 'nov': '11', 'dic': '12',
    'enero': '01', 'febrero': '02', 'marzo': '03', 'abril': '04',
    'mayo': '05', 'junio': '06', 'julio': '07', 'agosto': '08',
    'septiembre': '09', 'setiembre': '09', 'octubre': '10', 'noviembre': '11', 'diciembre': '12'
}

class SimpleHTMLTableParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.rows = []
        self.current_row = []
        self.current_cell = []
        self.in_cell = False

    def handle_starttag(self, tag, attrs):
        t = tag.lower()
        if t in ('td', 'th'):
            self.in_cell = True
            self.current_cell = []
        elif t == 'tr':
            self.current_row = []

    def handle_data(self, data):
        if self.in_cell:
            self.current_cell.append(data)

    def handle_endtag(self, tag):
        t = tag.lower()
        if t in ('td', 'th'):
            self.in_cell = False
            cell_text = ' '.join(''.join(self.current_cell).split())
            self.current_row.append(cell_text)
        elif t == 'tr':
            if any(c.strip() for c in self.current_row):
                self.rows.append(self.current_row)

def clean_amount(val, is_credit_card_charge: bool = False) -> float:
    if val is None or val == "":
        return 0.0
    if isinstance(val, (int, float)):
        amt = float(val)
        return -abs(amt) if is_credit_card_charge and amt > 0 else amt

    s = str(val).strip().replace('\xa0', ' ').replace('&nbsp;', ' ')

    # Normalize various unicode minus / dash symbols
    s = s.replace('\u2212', '-')  # Unicode minus (8722)
    s = s.replace('\u2013', '-')  # En dash
    s = s.replace('\u2014', '-')  # Em dash
    s = s.replace('\uff0d', '-')  # Fullwidth hyphen-minus
    s = s.replace('\ufe63', '-')  # Small hyphen-minus

    # Check for negative signs or parenthesis: (12.50) -> -12.50, or minus at the end: 12.50-
    is_negative = '-' in s or (s.startswith('(') and s.endswith(')')) or s.endswith('-')

    # Remove currency words (EUR, USD, etc.) and symbols (€, $, £, quotes, parens)
    s = re.sub(r'(?i)\b(eur|euros?|usd|dollars?|gbp|chf)\b', '', s)
    s = re.sub(r'[€$£\s"\'\(\)]', '', s)
    s = s.replace('-', '') # remove minus to format cleanly

    if not s:
        return 0.0

    # European format (1.250,50) vs US format (1,250.50)
    if ',' in s and '.' in s:
        if s.rfind(',') > s.rfind('.'):
            # European: 1.250,50 -> 1250.50
            s = s.replace('.', '').replace(',', '.')
        else:
            # US: 1,250.50 -> 1250.50
            s = s.replace(',', '')
    elif ',' in s:
        # 1250,50 -> 1250.50
        s = s.replace(',', '.')

    try:
        num = float(s)
        if is_negative:
            return -abs(num)
        if is_credit_card_charge:
            return -abs(num)
        return num
    except ValueError:
        return 0.0

def parse_date(val):
    if val is None or not str(val).strip():
        return None
    if isinstance(val, datetime):
        return val.strftime("%Y-%m-%d")

    s = str(val).strip().replace('\xa0', ' ')

    # Support Excel serial date numbers (e.g. 46038.0 or 46038)
    try:
        num = float(s)
        if 25000 <= num <= 70000:
            base_date = datetime(1899, 12, 30)
            return (base_date + timedelta(days=num)).strftime("%Y-%m-%d")
    except (ValueError, TypeError, OverflowError):
        pass

    # Remove pipe symbol or ISO 'T'
    if '|' in s:
        s = s.split('|')[0].strip()
    if 'T' in s and re.search(r'\d{4}-\d{2}-\d{2}T', s):
        s = s.split('T')[0].strip()

    # Remove time of day if present, e.g. " 14:30:00" or " 14:30"
    s = re.sub(r'\s+\d{1,2}:\d{2}(:\d{2})?(\s*(am|pm|AM|PM))?.*$', '', s).strip()

    # Normalize separators
    s = s.replace('.', '/').replace('-', '/')

    # Replace Spanish month names if any (e.g. 11/sep/2026 -> 11/09/2026)
    s_lower = s.lower()
    for m_name, m_num in SPANISH_MONTH_MAP.items():
        if m_name in s_lower:
            s_lower = re.sub(r'\b' + m_name + r'\b', m_num, s_lower)
            s = s_lower

    # If spaces remain (e.g. "11 09 2026")
    s = re.sub(r'\s+', '/', s)

    formats = ["%d/%m/%Y", "%Y/%m/%d", "%d/%m/%y"]
    for fmt in formats:
        try:
            return datetime.strptime(s, fmt).strftime("%Y-%m-%d")
        except ValueError:
            pass

    return None

def detect_account_from_content(full_text: str = "", default_account: str = "Santander (*3012)", filename: str = "") -> str:
    """
    Returns the account selected by the user, or auto-detects if 'Auto' or not explicitly specified.
    """
    if default_account and default_account != "Auto":
        return default_account

    sample = (filename + " " + full_text[:4000]).lower()
    if any(k in sample for k in ["myinvestor", "my investor", "fidelity msci world", "transferencia desde myinvestor"]):
        return ACCOUNT_MYINVESTOR
    if "n26" in sample:
        return ACCOUNT_N26
    if "iberia cards" in sample or "iberiacards" in sample:
        return ACCOUNT_IBERIA_CARDS
    if "american express" in sample or "amex" in sample:
        return ACCOUNT_AMEX
    return ACCOUNT_SANTANDER_MAIN

def parse_csv_content(content_str: str, default_account: str = "Santander (*3012)", filename: str = ""):
    """
    Parses CSV content from Santander, Iberia Cards, N26, Amex, or generic extracts.
    Robust against Windows/Mac line endings (\r\n, \r), unquoted fields, and diverse delimiters.
    """
    detected_account = detect_account_from_content(content_str, default_account, filename=filename)
    is_credit_card = "Iberia Cards" in detected_account or "American Express" in detected_account or "Santander Tarjeta" in detected_account or "Tarjeta" in detected_account

    # Normalize line breaks to standard Unix LF (\n) to prevent _csv.Error on \r carriage returns
    normalized_content = content_str.replace('\r\n', '\n').replace('\r', '\n')

    lines = [line for line in normalized_content.split('\n') if line.strip()]
    if not lines:
        return []

    # Detect delimiter (; or , or \t)
    sample_block = "\n".join(lines[:20])
    comma_count = sample_block.count(',')
    semicolon_count = sample_block.count(';')
    tab_count = sample_block.count('\t')

    if tab_count > comma_count and tab_count > semicolon_count:
        delimiter = '\t'
    elif semicolon_count > comma_count:
        delimiter = ';'
    else:
        delimiter = ','

    # Resilient CSV parser: try standard reader first, fallback to line-by-line if malformed quotes
    all_rows = []
    try:
        reader = csv.reader(io.StringIO(normalized_content, newline=''), delimiter=delimiter)
        for r in reader:
            if any(cell.strip() for cell in r):
                all_rows.append([cell.strip() for cell in r])
    except Exception as e:
        print(f"Standard csv.reader note: {e}. Switching to resilient line parser.")
        all_rows = []
        for line in lines:
            try:
                r = next(csv.reader([line], delimiter=delimiter))
                if any(c.strip() for c in r):
                    all_rows.append([c.strip() for c in r])
            except Exception:
                r = [c.strip().strip('"\'') for c in line.split(delimiter)]
                if any(r):
                    all_rows.append(r)

    if not all_rows:
        return []

    # Find the header row
    header_idx = -1
    date_col = -1
    desc_col = -1
    ref_col = -1
    amount_col = -1
    balance_col = -1
    currency_col = -1
    cargo_col = -1
    abono_col = -1

    for idx, row in enumerate(all_rows[:40]):
        row_str = " ".join(row).lower()
        has_date_term = any(w in row_str for w in ["fecha operación", "fecha operacion", "f. operacion", "transaction date", "booking date", "fecha", "date"])
        has_desc_term = any(w in row_str for w in ["comercio", "concepto", "descripción", "descripcion", "description", "movimiento", "detalle", "partner name", "partner", "payee", "beneficiario", "destinatario"])
        has_amount_term = any(w in row_str for w in ["importe", "amount", "cantidad", "saldo", "balance", "cargo", "abono"])

        if has_date_term and (has_desc_term or has_amount_term):
            header_idx = idx
            break

    if header_idx == -1:
        # Fallback: find any row with 'fecha' or 'date'
        for idx, row in enumerate(all_rows[:20]):
            row_str = " ".join(row).lower()
            if "fecha" in row_str or "date" in row_str:
                header_idx = idx
                break

    if header_idx == -1:
        header_idx = 0

    headers = [h.lower().strip() for h in all_rows[header_idx]]

    # Prioritized column matching:
    # 1. Date: First look strictly for Fecha de Operación (actual day the purchase/transfer occurred)
    for i, h in enumerate(headers):
        if any(w in h for w in ["fecha operación", "fecha operacion", "f. operacion", "f.operacion", "fecha transaccion", "transaction date", "operation date"]):
            date_col = i
            break
    if date_col == -1:
        for i, h in enumerate(headers):
            if "valor" not in h and "value" not in h and any(w in h for w in ["booking date", "fecha", "date"]):
                date_col = i
                break
    if date_col == -1:
        for i, h in enumerate(headers):
            if any(w in h for w in ["fecha", "date"]):
                date_col = i
                break

    # 2. First look for specific EUR amount columns before generic ones
    for i, h in enumerate(headers):
        if amount_col == -1 and any(w in h for w in ["importe euros", "amount (eur)", "importe neto", "importenet"]):
            amount_col = i

    orig_amount_col = -1
    orig_curr_col = -1

    for i, h in enumerate(headers):
        if desc_col == -1 and any(w in h for w in ["partner name", "partner", "destinatario", "ordenante", "comercio", "description", "concepto", "descripción", "descripcion", "detalle", "movimiento", "payee", "beneficiario"]):
            desc_col = i
        elif ref_col == -1 and any(w in h for w in ["payment reference", "referencia de pago", "verwendungszweck", "concepto 2", "referencia"]):
            ref_col = i
        elif balance_col == -1 and any(w in h for w in ["balance", "saldo", "saldo contable", "saldo disponible", "saldo actual"]):
            balance_col = i
        elif amount_col == -1 and any(w in h for w in ["importe", "amount", "cantidad", "cargo/abono"]):
            if "divisa" not in h and "foreign" not in h and "original" not in h:
                amount_col = i
        elif currency_col == -1 and any(w in h for w in ["currency", "divisa", "moneda"]):
            if "amount" not in h and "importe" not in h and "original" not in h:
                currency_col = i
        elif orig_amount_col == -1 and any(w in h for w in ["original amount", "importe original"]):
            orig_amount_col = i
        elif orig_curr_col == -1 and any(w in h for w in ["original currency", "divisa original"]):
            orig_curr_col = i

        if cargo_col == -1 and any(w in h for w in ["cargo", "cargos", "adeudo", "adeudos", "debe"]):
            cargo_col = i
        if abono_col == -1 and any(w in h for w in ["abono", "abonos", "ingreso", "ingresos", "haber"]):
            abono_col = i

    # Fallbacks if columns weren't matched by exact headers
    if date_col == -1: date_col = 0
    if desc_col == -1: desc_col = 2 if len(headers) > 2 else (1 if len(headers) > 1 else 0)
    if amount_col == -1 and cargo_col == -1 and abono_col == -1:
        amount_col = 3 if len(headers) > 3 else (2 if len(headers) > 2 else -1)

    category_col = -1
    for i, h in enumerate(headers):
        if category_col == -1 and any(w in h for w in ["categoría", "categoria", "category"]):
            category_col = i

    is_amex = detected_account == ACCOUNT_AMEX or "American Express" in detected_account or "AMEX" in detected_account.upper()

    # Detect header balance in rows before header_idx if not in column
    header_balance = None
    for row in all_rows[:header_idx]:
        for c_idx, cell in enumerate(row):
            if cell.strip().lower() in ["balance", "saldo", "saldo actual", "saldo disponible"]:
                if c_idx + 1 < len(row) and row[c_idx + 1].strip():
                    val = clean_amount(row[c_idx + 1])
                    if val != 0.0 or "0" in row[c_idx + 1]:
                        header_balance = val
                        break

    # Check if this file already has explicit signed amounts (e.g. Santander accounts/cards, N26 have '-' or '−')
    has_negative_amounts = False
    if amount_col >= 0:
        for r in all_rows[header_idx + 1:]:
            if amount_col < len(r):
                val_str = r[amount_col].strip()
                if any(c in val_str for c in ['-', '\u2212', '\u2013', '\u2014', '\uff0d', '\ufe63']) or (val_str.startswith('(') and val_str.endswith(')')):
                    has_negative_amounts = True
                    break

    # Only treat unsigned amounts as charges if the account is an unsigned card extract (like Iberia Cards statements)
    # AND the extract does not already have explicit negative signs.
    is_unsigned_card_statement = ("Iberia Cards" in detected_account) and not has_negative_amounts

    parsed_txs = []
    for row in all_rows[header_idx + 1:]:
        if len(row) <= max(date_col, desc_col):
            continue

        raw_date = row[date_col].strip()
        raw_desc = row[desc_col].strip() if desc_col < len(row) else ""
        raw_ref = row[ref_col].strip() if ref_col >= 0 and ref_col < len(row) else ""
        raw_orig_amt = row[orig_amount_col].strip() if orig_amount_col >= 0 and orig_amount_col < len(row) else ""
        raw_orig_curr = row[orig_curr_col].strip() if orig_curr_col >= 0 and orig_curr_col < len(row) else ""

        # Validate date
        parsed_date = parse_date(raw_date)
        if not parsed_date:
            continue

        # Clean multiline in strings (e.g. N26\n)
        raw_desc = raw_desc.replace('\n', ' ').strip()
        raw_ref = raw_ref.replace('\n', ' ').strip()

        # Enrich description with payment reference if available (very helpful for N26 and Amex)
        if not raw_desc and raw_ref:
            raw_desc = raw_ref
        elif raw_desc and raw_ref and raw_ref.lower() not in raw_desc.lower() and len(raw_ref) > 1:
            raw_desc = f"{raw_desc} - {raw_ref}"

        # Append original currency info if non-EUR
        if raw_orig_curr and raw_orig_curr.upper() != "EUR" and raw_orig_amt:
            raw_desc = f"{raw_desc} ({raw_orig_amt} {raw_orig_curr.upper()})"

        if not raw_desc:
            raw_desc = "Movimiento N26" if "N26" in detected_account else "Movimiento bancario"

        # Skip summary rows or empty rows
        if raw_desc.lower() in ["total a cargar", "total comercios", "total intereses", "total", "total cajeros", "deuda pendiente del mes anterior", "pagos y abonos", "cargos"]:
            continue

        raw_bank_category = row[category_col].strip() if category_col >= 0 and category_col < len(row) else ""

        if is_amex:
            # American Express statements use cardholder billing conventions:
            # - Positive amounts are charges/expenses (purchases, monthly fees) -> must be negative in personal finance
            # - Negative amounts are payments received or refunds -> must be positive in personal finance
            raw_amount = row[amount_col].strip() if amount_col >= 0 and amount_col < len(row) else "0"
            clean_num = clean_amount(raw_amount)
            amt = -clean_num
        else:
            # Check if this row is a credit card charge for unsigned extracts (Iberia Cards)
            is_card_charge = False
            if is_unsigned_card_statement:
                is_refund = any(kw in raw_desc.upper() for kw in ["ABONO", "DEVOLUCION", "REEMBOLSO"])
                is_card_charge = not is_refund

            if amount_col >= 0 and amount_col < len(row) and row[amount_col].strip():
                raw_amount = row[amount_col].strip()
                amt = clean_amount(raw_amount, is_credit_card_charge=is_card_charge)
            elif cargo_col >= 0 and cargo_col < len(row) and row[cargo_col].strip() and clean_amount(row[cargo_col].strip()) != 0:
                amt = -abs(clean_amount(row[cargo_col].strip()))
            elif abono_col >= 0 and abono_col < len(row) and row[abono_col].strip() and clean_amount(row[abono_col].strip()) != 0:
                amt = abs(clean_amount(row[abono_col].strip()))
            else:
                raw_amount = row[amount_col].strip() if amount_col >= 0 and amount_col < len(row) else "0"
                amt = clean_amount(raw_amount, is_credit_card_charge=is_card_charge)
        if not raw_desc or raw_desc.strip().lower() in ["null", "none", "movimiento bancario"]:
            if "myinvestor" in detected_account.lower():
                raw_desc = "Abono MyInvestor" if amt >= 0 else "Cargo MyInvestor"
            elif "n26" in detected_account.lower():
                raw_desc = "Movimiento N26"
            else:
                raw_desc = "Movimiento bancario"

        balance_val = None
        if balance_col >= 0 and balance_col < len(row) and row[balance_col].strip():
            balance_val = clean_amount(row[balance_col].strip())
        elif header_balance is not None and not parsed_txs:
            balance_val = header_balance

        currency = "EUR"
        if currency_col >= 0 and currency_col < len(row) and row[currency_col].strip():
            candidate_curr = row[currency_col].strip().upper()
            if not re.search(r'\d', candidate_curr) and len(candidate_curr) <= 5:
                currency = candidate_curr

        parsed_txs.append({
            "date": parsed_date,
            "account": detected_account,
            "description": raw_desc,
            "raw_description": raw_desc,
            "amount": amt,
            "balance": balance_val,
            "currency": currency,
            "bank_category": raw_bank_category
        })

    return parsed_txs

def extract_rows_from_bytes(file_bytes: bytes, filename: str = "") -> list:
    rows = []

    # 1. OpenXML .xlsx (starts with PK\x03\x04)
    if file_bytes.startswith(b'PK\x03\x04'):
        try:
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(file_bytes), data_only=True)
            ws = wb.active
            for r in ws.iter_rows(values_only=True):
                if any(c is not None and str(c).strip() != "" for c in r):
                    row_strs = []
                    for c in r:
                        if c is None:
                            row_strs.append("")
                        elif isinstance(c, datetime):
                            row_strs.append(c.strftime("%Y-%m-%d"))
                        elif isinstance(c, float) and c.is_integer():
                            row_strs.append(str(int(c)))
                        else:
                            row_strs.append(str(c).strip())
                    rows.append(row_strs)
            if rows:
                return rows
        except Exception as e:
            print("openpyxl extraction error:", e)

    # 2. Binary BIFF8 .xls (starts with OLE header \xd0\xcf\x11\xe0)
    if file_bytes.startswith(b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1'):
        try:
            import xlrd
            wb = xlrd.open_workbook(file_contents=file_bytes)
            ws = wb.sheet_by_index(0)
            for r_idx in range(ws.nrows):
                r_vals = ws.row_values(r_idx)
                if any(str(c).strip() != "" for c in r_vals):
                    row_strs = []
                    for c in r_vals:
                        if isinstance(c, float) and c.is_integer():
                            row_strs.append(str(int(c)))
                        else:
                            row_strs.append(str(c).strip())
                    rows.append(row_strs)
            if rows:
                return rows
        except Exception as e:
            print("xlrd extraction error:", e)

    # 3. HTML table disguised as .xls (Santander, BBVA, CaixaBank .xls exports)
    content_str = None
    for enc in ["utf-8-sig", "utf-8", "latin-1", "cp1252", "iso-8859-1"]:
        try:
            content_str = file_bytes.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if content_str is None:
        content_str = file_bytes.decode("utf-8", errors="ignore")

    content_lower = content_str.lower()
    if "<table" in content_lower and ("<tr" in content_lower or "<th" in content_lower or "<td" in content_lower):
        try:
            parser = SimpleHTMLTableParser()
            parser.feed(content_str)
            if parser.rows:
                return parser.rows
        except Exception as e:
            print("HTML table extraction error:", e)

    # 4. XML Spreadsheet 2003 (<Workbook ... <Row> ... <Cell>)
    if "<workbook" in content_lower and "<row" in content_lower:
        try:
            import xml.etree.ElementTree as ET
            xml_clean = re.sub(r' xmlns="[^"]+"', '', content_str, count=1)
            xml_clean = re.sub(r' xmlns:[a-zA-Z0-9]+="[^"]+"', '', xml_clean)
            xml_clean = re.sub(r' [a-zA-Z0-9]+:([a-zA-Z0-9]+)=', r' \1=', xml_clean)
            xml_clean = re.sub(r'</?[a-zA-Z0-9]+:', '<', xml_clean)
            root = ET.fromstring(xml_clean)
            for row_elem in root.iter("Row"):
                row = []
                for cell in row_elem.iter("Cell"):
                    data_elem = cell.find("Data")
                    val = data_elem.text if data_elem is not None and data_elem.text else ""
                    row.append(val.strip())
                if any(row):
                    rows.append(row)
            if rows:
                return rows
        except Exception as e:
            print("XML spreadsheet extraction error:", e)

    # 5. Plain CSV / TSV fallback
    try:
        lines = [line.strip() for line in content_str.splitlines() if line.strip()]
        for line in lines:
            delimiter = ';' if ';' in line else ('\t' if '\t' in line else ',')
            try:
                r = next(csv.reader([line], delimiter=delimiter))
                if any(c.strip() for c in r):
                    rows.append([c.strip() for c in r])
            except Exception:
                r = [c.strip().strip('"\'') for c in line.split(delimiter)]
                if any(r):
                    rows.append(r)
        if rows:
            return rows
    except Exception as e:
        print("CSV fallback extraction error:", e)

    return rows

def parse_excel_content(file_bytes: bytes, default_account: str = "Santander (*3012)", filename: str = ""):
    rows = extract_rows_from_bytes(file_bytes, filename=filename)
    if not rows:
        return []

    # Check for metadata balance in top rows
    header_balance = None
    for row in rows[:10]:
        for c_idx, cell in enumerate(row):
            if str(cell).strip().lower() in ["balance", "saldo", "saldo actual", "saldo disponible"]:
                if c_idx + 1 < len(row) and row[c_idx + 1].strip():
                    val = clean_amount(row[c_idx + 1])
                    if val != 0.0 or "0" in row[c_idx + 1]:
                        header_balance = val
                        break
        if header_balance is not None:
            break

    output = io.StringIO()
    writer = csv.writer(output, delimiter=';')
    writer.writerows(rows)
    txs = parse_csv_content(output.getvalue(), default_account=default_account, filename=filename)
    if header_balance is not None and txs:
        if txs[0].get("balance") is None:
            txs[0]["balance"] = header_balance
    return txs
